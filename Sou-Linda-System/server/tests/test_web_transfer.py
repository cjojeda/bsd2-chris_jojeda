"""HTTP file boundaries use the same transactional engine as notebook operations."""
from contextlib import closing
from copy import deepcopy
import base64
import csv
import io
import json
from pathlib import Path
import sqlite3
import sys

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from sadi_local import transfer
from sadi_local.web import create_app
from sadi_local import web_transfer


@pytest.fixture
def api(tmp_path):
    app = create_app(tmp_path / "business.sqlite3", allow_mobile=True,
                     report_addresses=["http://192.168.1.20:8765"])
    with TestClient(app, base_url="http://localhost:8765", client=("127.0.0.1", 50001)) as client:
        headers = {"X-SADI-CSRF": client.get("/api/session").json()["csrfToken"]}
        yield client, headers, tmp_path


def post(api, path, payload):
    client, headers, _ = api
    return client.post(path, headers=headers, json=payload)


def variant(api, sku="VAL-95-R", **changes):
    payload = {"name": "Valentina", "category": "Conjuntos", "sku": sku, "size": "95", "color": "Rojo",
               "cost": 10000, "retailPrice": 25000, "wholesalePrice": 20000, "stock": 10, "minStock": 2}
    payload.update(changes)
    response = post(api, "/api/variants", payload)
    assert response.status_code == 200, response.text
    return response.json()["id"]


def sale(api, variant_id):
    response = post(api, "/api/sales", {"date": "2026-10-01", "channel": "Instagram", "paymentMethod": "Transferencia",
        "items": [{"variantId": variant_id, "quantity": 2, "unitPrice": 25000}], "commandId": "transfer-sale"})
    assert response.status_code == 200, response.text
    return response.json()


def csv_upload(rows, file_name="productos.csv"):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter=";")
    writer.writeheader()
    writer.writerows(rows)
    return upload(("\ufeff" + stream.getvalue()).encode("utf-8"), file_name)


def upload(content, file_name):
    return {"fileName": file_name, "contentBase64": base64.b64encode(content).decode("ascii")}


def product_row(**changes):
    result = {"modelo": "Valentina", "categoria": "Conjuntos", "sku": "VAL-95-R", "talle": "95", "color": "Rojo",
              "costo": "100,00", "pvp": "250,00", "mayorista": "200,00", "stock": "10", "stock_minimo": "2"}
    result.update(changes)
    return result


def history_row(**changes):
    result = {"id_externo": "TREINTA-1", "fecha": "10/09/2026", "sku": "VAL-95-R", "cantidad": "2",
              "precio_unitario": "250,00", "costo_unitario": "100,00", "canal": "Instagram",
              "medio_pago": "Transferencia", "envio": "10,00", "packaging": "5,00", "descuento": "30,00",
              "comision_porcentaje": "3", "cliente": "", "cliente_id": "", "fecha_cobro": "11/09/2026"}
    result.update(changes)
    return result


def preview(api, file, kind="variants", **changes):
    return post(api, "/api/imports/preview", {**file, "kind": kind, "sourceAccount": "cuenta-negocio", **changes})


def apply(api, response, acknowledge=False):
    assert response.status_code == 200, response.text
    return post(api, "/api/imports/apply", {"previewId": response.json()["previewId"], "acknowledgeWarnings": acknowledge})


def test_json_and_sqlite_downloads_are_complete_branded_and_do_not_contain_auth_secrets(api):
    client, headers, tmp_path = api
    identifier = variant(api)
    sale(api, identifier)
    token = client.post("/api/mobile/pair", headers=headers).json()["token"]
    response = client.get("/api/backup")
    assert response.status_code == 200
    assert "sou-linda-backup-" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    assert token not in response.text and "csrfToken" not in response.text
    document = response.json()
    assert document["format"] == transfer.BACKUP_FORMAT
    tables = transfer.validate_backup(document)
    assert len(tables["orders"]) == 1 and len(tables["variants"]) == 1
    sqlite_copy = client.get("/api/backup/sqlite")
    assert sqlite_copy.status_code == 200 and sqlite_copy.content.startswith(b"SQLite format 3\x00")
    assert "sou-linda-base-" in sqlite_copy.headers["content-disposition"]
    target = tmp_path / "download.sqlite3"
    target.write_bytes(sqlite_copy.content)
    with closing(sqlite3.connect(target)) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("SELECT count(*) FROM orders").fetchone()[0] == 1
        assert connection.execute("SELECT on_hand FROM inventory").fetchone()[0] == 8
    assert not any(name.startswith("PySide6") for name in sys.modules)


def test_json_restore_replaces_all_data_and_preserves_automatic_previous_backup(api):
    client, _, tmp_path = api
    identifier = variant(api)
    sale(api, identifier)
    content = client.get("/api/backup").text
    validation = post(api, "/api/backup/validate", {"content": content})
    assert validation.status_code == 200
    assert validation.json()["counts"]["orders"] == 1
    variant(api, "EXTRA", name="Anterior")
    before = client.get("/api/snapshot").json()
    assert len(before["variants"]) == 2
    rejected = post(api, "/api/backup/restore", {"content": content, "confirmed": False})
    assert rejected.status_code == 400
    assert client.get("/api/snapshot").json() == before
    restored = post(api, "/api/backup/restore", {"content": content, "confirmed": True})
    assert restored.status_code == 200, restored.text
    assert len(client.get("/api/snapshot").json()["variants"]) == 1
    previous = Path(restored.json()["previousBackup"])
    assert previous.parent == tmp_path / "backups" and previous.exists()
    previous_document = json.loads(previous.read_text(encoding="utf-8"))
    assert len(transfer.validate_backup(previous_document)["variants"]) == 2


@pytest.mark.parametrize("content", ["{broken", '{"x": 1, "x": 2}', "[]", "null"])
def test_invalid_backup_never_changes_active_database(api, content):
    client, _, _ = api
    variant(api)
    before = client.get("/api/snapshot").json()
    for path in ("/api/backup/validate", "/api/backup/restore"):
        payload = {"content": content}
        if path.endswith("restore"):
            payload["confirmed"] = True
        assert post(api, path, payload).status_code == 400
        assert client.get("/api/snapshot").json() == before


def test_checksum_damage_or_resigned_stock_corruption_cannot_be_restored(api):
    client, _, _ = api
    variant(api)
    original = client.get("/api/backup").json()
    before = client.get("/api/snapshot").json()
    damaged = deepcopy(original)
    damaged["payload"]["tables"]["variants"][0]["retail_cents"] += 1
    assert post(api, "/api/backup/restore", {"content": json.dumps(damaged), "confirmed": True}).status_code == 400
    resigned = deepcopy(original)
    resigned["payload"]["tables"]["inventory"][0]["on_hand"] += 1
    resigned["sha256"] = transfer._hash({key: value for key, value in resigned.items() if key != "sha256"})
    assert post(api, "/api/backup/restore", {"content": json.dumps(resigned), "confirmed": True}).status_code == 400
    assert client.get("/api/snapshot").json() == before


def test_csv_column_mapping_preview_is_read_only_and_commit_is_idempotent(api):
    client, _, _ = api
    renamed = {"Articulo": "Valentina", "Codigo": "VAL-95-R", "Tamaño": "95", "Tono": "Rojo",
               "Compra": "100,00", "Publico": "250,00", "Mayor": "200,00", "Unidades": "10"}
    file = csv_upload([renamed])
    inspected = post(api, "/api/imports/inspect", file)
    assert inspected.status_code == 200
    assert inspected.json()["rowCount"] == 1 and inspected.json()["headers"] == list(renamed)
    mapping = dict(zip(("modelo", "sku", "talle", "color", "costo", "pvp", "mayorista", "stock"), renamed))
    result = preview(api, file, mapping=mapping)
    assert result.status_code == 200 and result.json()["ready"] is True
    assert result.json()["documentCount"] == 1
    assert client.get("/api/snapshot").json()["variants"] == []
    first, second = apply(api, result), apply(api, result)
    assert first.status_code == second.status_code == 200
    assert first.json() == second.json() and first.json()["count"] == 1
    assert len(client.get("/api/snapshot").json()["variants"]) == 1


def test_historical_rows_group_order_keep_current_stock_and_require_warning_acknowledgement(api):
    client, _, _ = api
    variant(api)
    result = preview(api, csv_upload([history_row(), history_row(cantidad="1")], "ventas.csv"), "sales")
    assert result.status_code == 200, result.text
    assert result.json()["ready"] and result.json()["documentCount"] == 1
    assert result.json()["warnings"]
    assert apply(api, result).status_code == 400
    assert client.get("/api/snapshot").json()["sales"] == []
    imported = apply(api, result, acknowledge=True)
    assert imported.status_code == 200, imported.text
    snapshot = client.get("/api/snapshot").json()
    assert snapshot["variants"][0]["stock"] == 10
    imported_sale = snapshot["sales"][0]
    assert imported_sale["total"] == 72000 and imported_sale["costTotal"] == 30000
    assert imported_sale["shippingCost"] == 1000 and imported_sale["packagingCost"] == 500
    assert imported_sale["source"] == "import"
    report = client.get("/api/reports/monthly?month=2026-09").json()
    assert report["net_profit_cents"] == 38340
    repeated = preview(api, csv_upload([history_row()], "otro.csv"), "sales")
    assert repeated.status_code == 200 and repeated.json()["ready"] is False
    assert any("ya fue importada" in error.lower() for error in repeated.json()["errors"])


def test_missing_historical_cost_yields_unknown_profit_instead_of_zero(api):
    client, _, _ = api
    variant(api)
    result = preview(api, csv_upload([history_row(costo_unitario="")], "ventas.csv"), "sales")
    assert result.json()["ready"] is True
    assert apply(api, result, acknowledge=True).status_code == 200
    report = client.get("/api/reports/monthly?month=2026-09").json()
    assert report["net_profit_cents"] is None and report["incomplete_cost_count"] == 1
    assert client.get("/api/snapshot").json()["sales"][0]["costTotal"] is None


def test_invalid_mapping_data_and_duplicate_headers_do_not_import(api):
    client, _, _ = api
    result = preview(api, csv_upload([product_row(stock="1,5")]))
    assert result.status_code == 200 and result.json()["ready"] is False
    assert apply(api, result).status_code == 400
    assert post(api, "/api/imports/inspect", upload(b"sku;SKU\nA;B\n", "duplicados.csv")).status_code == 400
    assert client.get("/api/snapshot").json()["variants"] == []


@pytest.mark.parametrize("file", [
    {"fileName": "macro.xlsm", "contentBase64": "YWJj"},
    {"fileName": "productos.csv", "contentBase64": "!!!!"},
    {"fileName": "productos.csv", "contentBase64": ""},
    {"fileName": "productos.csv", "contentBase64": "AA=="},
    {"fileName": "catalogo.xlsx", "contentBase64": "YWJj"},
])
def test_file_decode_and_type_errors_are_controlled(api, file):
    result = post(api, "/api/imports/inspect", file)
    assert result.status_code in (400, 422), result.text


def test_import_size_and_plan_expiration_are_bounded(api, monkeypatch):
    file = csv_upload([product_row()])
    monkeypatch.setattr(web_transfer, "MAX_IMPORT_BYTES", 4)
    assert post(api, "/api/imports/inspect", file).status_code == 400
    monkeypatch.setattr(web_transfer, "MAX_IMPORT_BYTES", 10 * 1024 * 1024)
    result = preview(api, file)
    assert result.status_code == 200
    monkeypatch.setattr(web_transfer, "PLAN_TTL_SECONDS", -1)
    assert apply(api, result).status_code == 410


def test_backup_decoded_utf8_size_is_bounded_without_touching_database(api, monkeypatch):
    client, _, _ = api
    before = client.get("/api/snapshot").json()
    monkeypatch.setattr(web_transfer, "MAX_BACKUP_BYTES", 8)
    result = post(api, "/api/backup/validate", {"content": "é" * 5})
    assert result.status_code == 400 and "50 MB" in result.text
    assert client.get("/api/snapshot").json() == before


def test_xlsx_inspection_and_numeric_export_reject_formulas(api):
    client, _, _ = api
    book = Workbook()
    sheet = book.active
    row = product_row()
    sheet.append(list(row))
    sheet.append(list(row.values()))
    output = io.BytesIO()
    book.save(output)
    file = upload(output.getvalue(), "catalogo.xlsx")
    result = preview(api, file)
    assert apply(api, result).status_code == 200
    response = client.get("/api/exports?kind=variants&format=xlsx")
    assert response.status_code == 200 and "sou-linda-variants.xlsx" in response.headers["content-disposition"]
    exported = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert exported.active.title == "Sou Linda System"
        data = list(exported.active.values)
        assert data[1][5] == 100 and data[1][6] == 250
    finally:
        exported.close()
    sheet["A2"] = '=HYPERLINK("https://example.invalid","click")'
    output = io.BytesIO()
    book.save(output)
    book.close()
    rejected = post(api, "/api/imports/inspect", upload(output.getvalue(), "con-formula.xlsx"))
    assert rejected.status_code == 400 and "fórmulas" in rejected.text


def test_csv_exports_escape_formula_like_model_names_and_templates_have_canonical_columns(api):
    client, _, _ = api
    variant(api, name="=SUM(A1:A9)")
    response = client.get("/api/exports?kind=variants&format=csv")
    assert response.status_code == 200 and "'=SUM(A1:A9)" in response.text
    template = client.get("/api/templates/sales")
    assert template.status_code == 200
    assert template.content.decode("utf-8-sig").strip().split(";") == list(transfer.SALE_COLUMNS)


def test_restore_invalidates_preview_and_clear_demo_cannot_destroy_real_data(api):
    client, _, tmp_path = api
    result = preview(api, csv_upload([product_row()]))
    original = client.get("/api/backup").text
    assert post(api, "/api/backup/restore", {"content": original, "confirmed": True}).status_code == 200
    assert apply(api, result).status_code == 410
    variant(api)
    before = client.get("/api/snapshot").json()
    assert post(api, "/api/demo/clear", {"confirmation": "ELIMINAR"}).status_code == 400
    assert client.get("/api/snapshot").json() == before


def test_clear_demo_requires_literal_confirmation_and_creates_previous_backup(api):
    client, _, _ = api
    assert post(api, "/api/demo", {}).status_code == 200
    before = client.get("/api/snapshot").json()
    assert post(api, "/api/demo/clear", {"confirmation": "eliminar"}).status_code == 400
    assert client.get("/api/snapshot").json() == before
    result = post(api, "/api/demo/clear", {"confirmation": "ELIMINAR"})
    assert result.status_code == 200, result.text
    previous = json.loads(Path(result.json()["previousBackup"]).read_text(encoding="utf-8"))
    assert len(transfer.validate_backup(previous)["orders"]) >= 20
    snapshot = client.get("/api/snapshot").json()
    assert snapshot["products"] == snapshot["sales"] == snapshot["customers"] == []


def test_lan_mobile_credentials_never_authorize_file_routes(api):
    client, headers, _ = api
    token = client.post("/api/mobile/pair", headers=headers).json()["token"]
    client._transport.client = ("192.168.1.50", 50002)
    forwarded = headers | {"Authorization": "Bearer " + token, "X-Forwarded-For": "127.0.0.1"}
    for path in ("/api/backup", "/api/backup/sqlite", "/api/exports?kind=variants&format=csv", "/api/templates/variants"):
        assert client.get(path, headers=forwarded).status_code == 403
    for path in ("/api/backup/validate", "/api/backup/restore", "/api/imports/inspect", "/api/imports/preview", "/api/imports/apply", "/api/demo/clear"):
        assert client.post(path, headers=forwarded, json={}).status_code == 403
