"""File-boundary tests: a backup must preserve business invariants, not just JSON."""
from contextlib import closing
from copy import deepcopy
import csv
from datetime import date
import io
import json
from pathlib import Path
import sqlite3

import pytest

from sadi_local.db import Database, TABLES
from sadi_local.demo import seed_demo
from sadi_local.reports import monthly_report
from sadi_local.services import SadiService
from sadi_local import transfer


@pytest.fixture
def service(tmp_path):
    db = Database(tmp_path / "active.sqlite3")
    yield SadiService(db)
    db.close()


def variant(service, sku="VAL-95-ROJ", **changes):
    data = dict(model="Valentina", category="Conjuntos", sku=sku, size="95", color="Rojo",
                retail_cents=25000, wholesale_cents=18000, min_stock=1,
                initial_qty=15, initial_unit_cost_cents=10000)
    data.update(changes)
    return service.save_variant(data)


def sale(service, variant_id, **changes):
    options = dict(business_date="2026-10-01", commission_bps=300, shipping_cents=500,
                   packaging_cents=200, discount_kind="percent", discount_value=1000)
    options.update(changes)
    return service.create_order([dict(variant_id=variant_id, quantity=2, unit_price_cents=25000)], "Instagram", **options)


def resign(document):
    unsigned = {key: value for key, value in document.items() if key != "sha256"}
    document["counts"] = {name: len(document["payload"]["tables"][name]) for name in TABLES}
    unsigned["counts"] = document["counts"]
    document["sha256"] = transfer._hash(unsigned)
    return document


def write_json(path, document):
    path.write_text(json.dumps(document, ensure_ascii=False), encoding="utf-8")
    return path


def csv_file(path, rows, fields=None):
    fields = fields or list(rows[0])
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, delimiter=";")
    writer.writeheader()
    writer.writerows(rows)
    path.write_text("\ufeff" + stream.getvalue(), encoding="utf-8")
    return path


def historic(sku="VAL-95-ROJ", **changes):
    row = dict(id_externo="T-1", fecha="2026-09-10", sku=sku, cantidad="2", precio_unitario="250,00",
               costo_unitario="100,00", canal="Instagram", medio_pago="Transferencia", envio="10,00",
               packaging="5,00", descuento="30,00", comision_porcentaje="3", cliente="", cliente_id="", fecha_cobro="2026-09-11")
    row.update(changes)
    return row


def test_json_roundtrip_restores_all_tables_and_preserves_previous_database(service, tmp_path):
    v = variant(service)
    sale(service, v)
    reserved = service.create_order([dict(variant_id=v, quantity=1, unit_price_cents=25000)], "WhatsApp", status="reserved")
    service.db.save_draft("pos", {"cart": [{"variant_id": v, "quantity": 2}], "channel": "TikTok"})
    path = transfer.export_json(service.db, tmp_path / "backup.json")
    expected = json.loads(path.read_text(encoding="utf-8"))["payload"]["tables"]
    other = Database(tmp_path / "other.sqlite3")
    try:
        old = SadiService(other)
        variant(old, sku="OLD", model="Anterior")
        before = other.snapshot()
        result = transfer.restore_json(other, path)
        assert other.snapshot() == expected
        assert other.load_draft("pos")["channel"] == "TikTok"
        assert old.get_order(reserved)["status"] == "reserved"
        previous = json.loads(Path(result["previous_backup"]).read_text(encoding="utf-8"))
        assert previous["payload"]["tables"] == before
        assert set(result["counts"]) == set(TABLES)
        transfer.validate_database(other)
    finally:
        other.close()


def test_backup_from_complete_demo_passes_business_validation(service):
    seed_demo(service)
    tables = transfer.validate_backup(transfer.make_backup(service.db))
    assert len(tables["orders"]) > 10
    assert any(row["state"] == "active" for row in tables["reservations"])


def test_restore_rejects_checksum_damage_without_touching_active_data(service, tmp_path):
    variant(service)
    before = service.db.snapshot()
    document = transfer.make_backup(service.db)
    document["payload"]["tables"]["variants"][0]["retail_cents"] += 1
    with pytest.raises(ValueError, match="checksum"):
        transfer.restore_json(service.db, write_json(tmp_path / "damaged.json", document))
    assert service.db.snapshot() == before


@pytest.mark.parametrize("corruption,match", [
    ("cash_amount", "caja"),
    ("cash_method", "caja"),
    ("cash_missing", "caja"),
    ("cash_duplicate", "caja"),
    ("commission_amount", "costos|comisiones"),
    ("commission_date", "costos|comisiones"),
    ("shipping_amount", "costos|comisiones"),
    ("expense_amount", "costos|comisiones"),
])
def test_resigned_backup_cannot_change_cash_or_cost_ledgers(service, corruption, match):
    v = variant(service)
    sale(service, v)
    service.save_expense("2026-10-01", "Publicidad", "Campaña", 3000, "Transferencia")
    document = transfer.make_backup(service.db)
    tables = document["payload"]["tables"]
    if corruption.startswith("cash"):
        row = next(row for row in tables["cash_entries"] if row["kind"] == "receipt")
        if corruption == "cash_amount":
            row["delta_cents"] += 1
        elif corruption == "cash_method":
            row["method"] = "Transferencia"
        elif corruption == "cash_missing":
            tables["cash_entries"].remove(row)
        else:
            duplicate = dict(row, id="duplicated-ledger-entry", sequence=max(r["sequence"] for r in tables["cash_entries"]) + 1)
            tables["cash_entries"].append(duplicate)
    else:
        bucket = {"commission_amount": "commission", "commission_date": "commission", "shipping_amount": "shipping", "expense_amount": "operating_expense"}[corruption]
        row = next(row for row in tables["economic_entries"] if row["bucket"] == bucket)
        if corruption == "commission_date":
            row["business_date"] = "2026-09-01"
        else:
            row["amount_cents"] += 1
    with pytest.raises(ValueError, match=match):
        transfer.validate_backup(resign(document))


@pytest.mark.parametrize("corruption", ["missing_fk", "stock_balance", "invalid_date", "float_money", "unknown_column", "missing_table"])
def test_resigned_malformed_database_is_rejected(service, corruption):
    v = variant(service)
    sale(service, v)
    document = transfer.make_backup(service.db)
    tables = document["payload"]["tables"]
    if corruption == "missing_fk":
        tables["order_items"][0]["variant_id"] = "not-a-variant"
    elif corruption == "stock_balance":
        tables["inventory"][0]["on_hand"] += 1
    elif corruption == "invalid_date":
        tables["payments"][0]["business_date"] = "2026-02-30"
    elif corruption == "float_money":
        tables["order_items"][0]["unit_price_cents"] = 25000.0
    elif corruption == "unknown_column":
        tables["variants"][0]["unknown"] = "discarding unknown data is unsafe"
    else:
        del tables["commands"]
        document["sha256"] = transfer._hash({key: value for key, value in document.items() if key != "sha256"})
        with pytest.raises(ValueError):
            transfer.validate_backup(document)
        return
    with pytest.raises(ValueError):
        transfer.validate_backup(resign(document))


def test_restore_rolls_back_when_storage_fails_after_replacement_started(service, tmp_path, monkeypatch):
    variant(service)
    before = service.db.snapshot()
    source = Database(":memory:")
    try:
        variant(SadiService(source), sku="NEW", model="Otra base")
        path = write_json(tmp_path / "valid.json", transfer.make_backup(source))
    finally:
        source.close()
    real_insert = transfer._insert_tables

    def fail_after_insert(db, tables):
        real_insert(db, tables)
        if db is service.db:
            raise sqlite3.OperationalError("simulated disk full")

    monkeypatch.setattr(transfer, "_insert_tables", fail_after_insert)
    with pytest.raises(ValueError, match="conservan"):
        transfer.restore_json(service.db, path)
    assert service.db.snapshot() == before
    assert list((Path(service.db.path).parent / "backups").glob("antes-de-restaurar-*.json"))


def test_restore_rejects_nonfinite_numbers_hidden_inside_draft_json(service):
    service.db.save_draft("pos", {"cart": []})
    document = transfer.make_backup(service.db)
    document["payload"]["tables"]["drafts"][0]["payload"] = '{"cart": [], "amount": NaN}'
    with pytest.raises(ValueError, match="borrador"):
        transfer.validate_backup(resign(document))


def test_sqlite_backup_includes_wal_and_closes_handle_before_windows_replace(service, tmp_path, monkeypatch):
    v = variant(service)
    sale(service, v)
    opened = []
    real_connect = sqlite3.connect
    real_replace = transfer.os.replace
    checked = []

    def connect(*args, **kwargs):
        target = real_connect(*args, **kwargs)
        opened.append(target)
        return target

    def replace(source, destination):
        assert opened
        with pytest.raises(sqlite3.ProgrammingError, match="closed"):
            opened[-1].execute("SELECT 1")
        checked.append(True)
        return real_replace(source, destination)

    monkeypatch.setattr(transfer.sqlite3, "connect", connect)
    monkeypatch.setattr(transfer.os, "replace", replace)
    path = transfer.export_sqlite(service.db, tmp_path / "backup.sqlite3")
    assert checked == [True]
    with closing(real_connect(path)) as saved:
        assert saved.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert saved.execute("SELECT COUNT(*) FROM orders").fetchone()[0] == 1
        assert saved.execute("SELECT on_hand FROM inventory WHERE variant_id=?", (v,)).fetchone()[0] == 13


@pytest.mark.parametrize("suffix", ["", "-wal", "-shm", "-journal"])
@pytest.mark.parametrize("operation", ["json", "sqlite", "table"])
def test_exports_never_overwrite_live_database_or_sqlite_sidecars(service, suffix, operation):
    variant(service)
    before = service.db.snapshot()
    target = Path(str(service.db.path) + suffix)
    with pytest.raises(ValueError, match="base abierta"):
        if operation == "json":
            transfer.export_json(service.db, target)
        elif operation == "sqlite":
            transfer.export_sqlite(service.db, target)
        else:
            transfer.export_table(service, target, "variants")
    assert service.db.snapshot() == before


def test_csv_parser_preserves_quoted_delimiters_quotes_multiline_and_bom(tmp_path):
    path = tmp_path / "quoted.csv"
    path.write_text('\ufeffsku;modelo\r\nAAA;"Conjunto; ""Seda""\nEdición especial"\r\n', encoding="utf-8")
    headers, rows = transfer.read_import_file(path)
    assert headers == ["sku", "modelo"]
    assert rows == [{"sku": "AAA", "modelo": 'Conjunto; "Seda"\nEdición especial'}]


@pytest.mark.parametrize("text", ["sku;SKU\nA;B", "sku;;modelo\nA;B;C", 'sku;modelo\nA;"comillas sin cerrar'])
def test_csv_rejects_ambiguous_headers_and_malformed_quotes(tmp_path, text):
    path = tmp_path / "bad.csv"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(ValueError):
        transfer.read_import_file(path)


def test_xlsx_parser_reads_actual_workbook_numbers_dates_and_rejects_formulas(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["sku", "fecha", "precio"])
    sheet.append(["AAA", date(2026, 9, 10), 1250.5])
    book.save(path)
    book.close()
    headers, rows = transfer.read_import_file(path)
    assert headers == ["sku", "fecha", "precio"]
    assert rows == [{"sku": "AAA", "fecha": "2026-09-10", "precio": "1250.5"}]
    book = Workbook()
    book.active.append(["sku", "precio"])
    book.active.append(["AAA", "=SUM(1,2)"])
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="fórmulas"):
        transfer.read_import_file(path)


def test_mapped_variant_import_uses_opening_stock_and_exact_money(service, tmp_path):
    path = csv_file(tmp_path / "catalog.csv", [dict(Nombre="Luna", SKU="LUN-95", Talle="95", Color="Marfil", Costo="100,25", Venta="250,50", Mayor="180", Unidades="3")])
    mapping = dict(modelo="Nombre", sku="SKU", talle="Talle", color="Color", costo="Costo", pvp="Venta", mayorista="Mayor", stock="Unidades")
    preview = transfer.preview_import(service, path, "variants", mapping)
    assert preview.ready
    assert not service.list_variants()
    assert transfer.apply_import(service, preview)["count"] == 1
    v = service.list_variants()[0]
    assert (v["model"], v["on_hand"], v["value_cents"], v["retail_cents"]) == ("Luna", 3, 30075, 25050)
    transfer.validate_database(service.db)


def test_catalog_preview_detects_duplicate_sku_and_normalized_variant(service, tmp_path):
    row = dict(modelo="Luna", categoria="Conjuntos", sku="LUN-95", talle="95", color="Marfil", costo="100", pvp="250", mayorista="180", stock="3", stock_minimo="1")
    path = csv_file(tmp_path / "catalog.csv", [row, dict(row, sku="OTHER", modelo="Lúna", color="MARFIL")])
    preview = transfer.preview_import(service, path, "variants")
    assert not preview.ready
    assert any("combinación" in error for error in preview.errors)
    with pytest.raises(ValueError):
        transfer.apply_import(service, preview)
    assert not service.list_variants()


def test_grouped_historical_items_share_one_discount_fee_and_cost_without_consuming_stock(service, tmp_path):
    first = variant(service)
    second = variant(service, sku="VAL-100-NEG", size="100", color="Negro")
    stock_before = service.db.snapshot()["inventory"]
    rows = [historic(), historic("VAL-100-NEG", cantidad="1", precio_unitario="300", costo_unitario="120")]
    path = csv_file(tmp_path / "history.csv", rows)
    preview = transfer.preview_import(service, path, "sales")
    assert preview.ready and len(preview.records) == 1
    with pytest.raises(ValueError, match="advertencias"):
        transfer.apply_import(service, preview)
    result = transfer.apply_import(service, preview, True)
    order = service.get_order(result["ids"][0])
    assert len(order["items"]) == 2
    assert {item["variant_id"] for item in order["items"]} == {first, second}
    assert order["subtotal_cents"] == 80000
    assert order["discount_cents"] == 3000
    assert order["total_cents"] == 77000
    assert order["cogs_cents"] == 32000
    assert order["commission_cents"] == 2310
    assert order["shipping_cents"] == 1000 and order["packaging_cents"] == 500
    assert order["net_profit_cents"] == 41190
    assert len(order["payments"]) == 1
    assert service.db.snapshot()["inventory"] == stock_before
    assert not service.c.execute("SELECT 1 FROM inventory_movements WHERE order_id=?", (order["id"],)).fetchone()
    assert service.c.execute("SELECT SUM(delta_cents) FROM cash_entries WHERE order_id=?", (order["id"],)).fetchone()[0] == 74690
    report = monthly_report(service.db, "2026-09")
    assert report["sales_count"] == 1 and report["net_profit_cents"] == 41190
    transfer.validate_backup(transfer.make_backup(service.db))


def test_historical_reimport_is_rejected_across_files_for_same_account(service, tmp_path):
    variant(service)
    first = csv_file(tmp_path / "first.csv", [historic()])
    transfer.apply_import(service, transfer.preview_import(service, first, "sales", source_account="mi-cuenta"), True)
    second = csv_file(tmp_path / "different-export.csv", [historic()])
    preview = transfer.preview_import(service, second, "sales", source_account="mi-cuenta")
    assert not preview.ready
    assert any("ya fue importada" in error for error in preview.errors)
    assert len(service.list_orders()) == 1
    # Distinct source accounts may legitimately use the same sale identifier.
    other = transfer.preview_import(service, second, "sales", source_account="otra-cuenta")
    assert other.ready
    transfer.apply_import(service, other, True)
    assert len(service.list_orders()) == 2


def test_external_identity_delimiters_cannot_collide(service, tmp_path):
    variant(service)
    left = csv_file(tmp_path / "one.csv", [historic(id_externo="c")])
    right = csv_file(tmp_path / "two.csv", [historic(id_externo="b:c")])
    transfer.apply_import(service, transfer.preview_import(service, left, "sales", source_account="a:b"), True)
    other = transfer.preview_import(service, right, "sales", source_account="a")
    assert other.ready
    transfer.apply_import(service, other, True)
    assert len(service.list_orders()) == 2


def test_unknown_historical_cost_remains_null_and_profit_is_incomplete(service, tmp_path):
    variant(service)
    path = csv_file(tmp_path / "unknown.csv", [historic(costo_unitario="")])
    preview = transfer.preview_import(service, path, "sales")
    assert preview.ready and any("falta costo histórico" in warning for warning in preview.warnings)
    result = transfer.apply_import(service, preview, True)
    order = service.get_order(result["ids"][0])
    assert order["cogs_cents"] is None and order["items"][0]["cogs_cents"] is None
    report = monthly_report(service.db, "2026-09")
    assert report["net_profit_cents"] is None and report["incomplete_cost_count"] == 1
    transfer.validate_backup(transfer.make_backup(service.db))


def test_mixed_known_and_missing_cost_preserves_known_portion_and_later_return(service, tmp_path):
    variant(service)
    variant(service, sku="VAL-100", size="100")
    path = csv_file(tmp_path / "mixed.csv", [historic(), historic("VAL-100", costo_unitario="", cantidad="1")])
    result = transfer.apply_import(service, transfer.preview_import(service, path, "sales"), True)
    order = service.get_order(result["ids"][0])
    known = next(item for item in order["items"] if item["cogs_cents"] is not None)
    assert service.c.execute("SELECT SUM(amount_cents) FROM economic_entries WHERE order_id=? AND bucket='cogs'", (order["id"],)).fetchone()[0] == known["cogs_cents"]
    service.return_order(order["id"], {known["id"]: 1}, business_date="2026-10-02")
    transfer.validate_backup(transfer.make_backup(service.db))


def test_missing_payment_method_is_not_invented(service, tmp_path):
    variant(service)
    path = csv_file(tmp_path / "missing-payment.csv", [historic(medio_pago="")])
    preview = transfer.preview_import(service, path, "sales")
    assert not preview.ready
    assert any("No se inventan cobros" in error for error in preview.errors)
    assert not service.list_orders()


def test_historical_credit_requires_client_and_does_not_invent_cash(service, tmp_path):
    variant(service)
    path = csv_file(tmp_path / "credit.csv", [historic(medio_pago="Crédito", cliente="Lucía")])
    result = transfer.apply_import(service, transfer.preview_import(service, path, "sales"), True)
    order = service.get_order(result["ids"][0])
    assert order["customer_name"] == "Lucía"
    assert order["balance_cents"] == 47000
    assert not order["payments"]
    assert not service.c.execute("SELECT 1 FROM cash_entries WHERE order_id=?", (order["id"],)).fetchone()


def test_import_rollback_removes_earlier_documents_when_later_write_fails(service, tmp_path, monkeypatch):
    variant(service)
    path = csv_file(tmp_path / "history.csv", [historic(id_externo="ONE"), historic(id_externo="TWO")])
    preview = transfer.preview_import(service, path, "sales")
    before = service.db.snapshot()
    real_insert = transfer._import_sale
    calls = []

    def fail_second(service_arg, group):
        calls.append(group["external_key"])
        if len(calls) == 2:
            raise sqlite3.OperationalError("simulated database failure")
        return real_insert(service_arg, group)

    monkeypatch.setattr(transfer, "_import_sale", fail_second)
    with pytest.raises(sqlite3.OperationalError):
        transfer.apply_import(service, preview, True)
    assert service.db.snapshot() == before


def test_import_revalidates_current_state_and_rejects_mutated_preview(service, tmp_path):
    variant(service)
    path = csv_file(tmp_path / "history.csv", [historic()])
    preview = transfer.preview_import(service, path, "sales")
    preview.raw_rows[0]["precio_unitario"] = "1"
    with pytest.raises(ValueError, match="vista previa cambió"):
        transfer.apply_import(service, preview, True)
    first = transfer.preview_import(service, path, "sales")
    stale = transfer.preview_import(service, path, "sales")
    transfer.apply_import(service, first, True)
    with pytest.raises(ValueError, match="ya fue importada"):
        transfer.apply_import(service, stale, True)
    assert len(service.list_orders()) == 1


def test_csv_export_escapes_formula_text_and_xlsx_keeps_it_as_literal_string(service, tmp_path):
    from openpyxl import load_workbook
    variant(service, model="=HYPERLINK(\"bad\")", sku="@SKU")
    path = transfer.export_table(service, tmp_path / "variants.csv", "variants")
    rows = list(csv.DictReader(io.StringIO(path.read_text(encoding="utf-8-sig")), delimiter=";"))
    assert rows[0]["modelo"].startswith("'=")
    assert rows[0]["sku"].startswith("'@")
    excel = transfer.export_table(service, tmp_path / "variants.xlsx", "variants")
    book = load_workbook(excel, read_only=True, data_only=False)
    try:
        cell = book.active["A2"]
        assert cell.value == '=HYPERLINK("bad")' and cell.data_type == "s"
    finally:
        book.close()


def test_exports_keep_negative_credit_numeric_and_exact(service, tmp_path):
    from openpyxl import load_workbook
    v = variant(service)
    order_id = sale(service, v, commission_bps=0)
    service.cancel_order(order_id, confirm_return=True, business_date="2026-10-02")
    target = transfer.export_table(service, tmp_path / "sales.csv", "sales")
    rows = list(csv.DictReader(io.StringIO(target.read_text(encoding="utf-8-sig")), delimiter=";"))
    assert rows[0]["saldo_pedido"] == "-450.00"
    target = transfer.export_table(service, tmp_path / "sales.xlsx", "sales")
    book = load_workbook(target, read_only=True, data_only=False)
    try:
        cell = book.active["O2"]
        assert cell.data_type == "n" and cell.value == -450
    finally:
        book.close()


def test_later_settlement_of_accrued_costs_is_valid_in_backup(service):
    v = variant(service)
    o = sale(service, v, direct_costs_paid=False)
    expense = service.save_expense("2026-10-01", "Publicidad", "Pendiente", 1500, "Transferencia", paid=False)
    service.pay_expense(expense, "QR", business_date="2026-10-03")
    service.pay_direct_costs(o, "Transferencia", business_date="2026-10-04")
    transfer.validate_backup(transfer.make_backup(service.db))


def test_refunds_and_inventory_loss_roundtrip_preserve_periods(service):
    v = variant(service)
    o = sale(service, v)
    service.return_order(o, business_date="2026-11-02")
    service.record_refund(o, 45000, "Transferencia", business_date="2026-11-03")
    service.adjust_stock(v, -1, "Prenda dañada", business_date="2026-11-04")
    transfer.validate_backup(transfer.make_backup(service.db))
    assert monthly_report(service.db, "2026-10")["revenue_cents"] == 45000
    assert monthly_report(service.db, "2026-11")["revenue_cents"] == -45000
