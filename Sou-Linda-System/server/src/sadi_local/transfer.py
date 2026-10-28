"""Versioned backups and deliberate CSV/Excel migration for Sou Linda System.

Files are untrusted input. Restores are validated in a disposable database before
the active database is modified; historical sales never consume today's stock.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from contextlib import closing
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
import csv
import hashlib
import io
import json
import os
import re
from pathlib import Path
import sqlite3
import tempfile
import unicodedata
from uuid import uuid4

from . import __version__
from .constants import CHANNELS, MAX_CENTS, PAYMENT_METHODS
from .db import Database, SCHEMA_VERSION, TABLES, utc_now
from .money import integer, parse_money, percent_bps
from .currency import currency_code, money_quantum, round_money_ratio, allocate_money, validate_amount

BACKUP_FORMAT = "sadi-desktop-backup"
BACKUP_VERSION = 1
MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_BACKUP_BYTES = 50 * 1024 * 1024
MAX_ROWS = 20_000

VARIANT_COLUMNS = ("modelo", "categoria", "sku", "talle", "color", "costo", "pvp", "mayorista", "stock", "stock_minimo")
SALE_COLUMNS = ("id_externo", "fecha", "sku", "cantidad", "precio_unitario", "costo_unitario", "canal", "medio_pago", "envio", "packaging", "descuento", "comision_porcentaje", "cliente", "cliente_id", "fecha_cobro")
REQUIRED_COLUMNS = {
    "variants": {"modelo", "sku", "talle", "color", "costo", "pvp", "mayorista", "stock"},
    "sales": {"id_externo", "fecha", "sku", "cantidad", "precio_unitario", "canal", "medio_pago"},
}


def _canonical(value) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _hash(value) -> str:
    return hashlib.sha256(_canonical(value)).hexdigest()


def _normal(value) -> str:
    return " ".join("".join(c for c in unicodedata.normalize("NFKD", str(value).strip()) if not unicodedata.combining(c)).casefold().split())


def _day(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    try:
        result = datetime.strptime(text, "%d/%m/%Y").date() if "/" in text else date.fromisoformat(text)
    except (TypeError, ValueError):
        raise ValueError(f"Fecha inválida: {text}. Usá AAAA-MM-DD o DD/MM/AAAA.") from None
    if "/" not in text and result.isoformat() != text:
        raise ValueError(f"Fecha inválida: {text}.")
    return result.isoformat()


def _atomic_file(path: Path, content: bytes) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp-" + uuid4().hex)
    try:
        with temporary.open("xb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return path


def _protect_database_path(db: Database, path: Path):
    if db.path != ":memory:":
        active = Path(db.path).resolve()
        reserved = {active, Path(str(active) + "-wal"), Path(str(active) + "-shm"), Path(str(active) + "-journal")}
        if path.resolve() in reserved:
            raise ValueError("Elegí otro archivo: no se puede sobrescribir la base abierta ni sus archivos auxiliares.")


def make_backup(db: Database) -> dict:
    tables = db.snapshot()
    document = {
        "format": BACKUP_FORMAT,
        "format_version": BACKUP_VERSION,
        "schema_version": SCHEMA_VERSION,
        "app_version": __version__,
        "created_at": utc_now(),
        "counts": {name: len(tables[name]) for name in TABLES},
        "payload": {"tables": tables},
    }
    return {**document, "sha256": _hash(document)}


def export_json(db: Database, path: str | Path) -> Path:
    path = Path(path)
    _protect_database_path(db, path)
    result = _atomic_file(path, json.dumps(make_backup(db), ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8"))
    db.set_meta("last_backup_at", utc_now())
    return result


def export_sqlite(db: Database, path: str | Path) -> Path:
    """SQLite's backup API includes the committed WAL, unlike copying the file."""
    path = Path(path)
    _protect_database_path(db, path)
    if db.connection.in_transaction:
        raise ValueError("Terminá la operación actual antes de respaldar SQLite.")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp-" + uuid4().hex)
    try:
        with closing(sqlite3.connect(temporary)) as target:
            db.connection.backup(target)
            if target.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise ValueError("El respaldo SQLite no superó el control de integridad.")
            if target.execute("PRAGMA foreign_key_check").fetchone() is not None:
                raise ValueError("El respaldo contiene referencias inválidas.")
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    db.set_meta("last_backup_at", utc_now())
    return path


def _no_duplicate_keys(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"El JSON repite la clave {key!r}.")
        result[key] = value
    return result


def _load_backup(path: str | Path) -> dict:
    path = Path(path)
    if path.stat().st_size > MAX_BACKUP_BYTES:
        raise ValueError("El respaldo supera el límite de 50 MB.")
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=_no_duplicate_keys,
                          parse_constant=lambda value: (_ for _ in ()).throw(ValueError(f"Número JSON inválido: {value}")))
    except (UnicodeError, json.JSONDecodeError) as error:
        raise ValueError("El archivo no es un respaldo JSON legible.") from error


def _validate_envelope(document):
    keys = {"format", "format_version", "schema_version", "app_version", "created_at", "counts", "payload", "sha256"}
    if not isinstance(document, dict) or set(document) != keys:
        raise ValueError("El archivo no tiene el formato completo de respaldo Sou Linda System.")
    if document["format"] != BACKUP_FORMAT:
        raise ValueError("Este archivo no usa el formato de respaldo compatible con Sou Linda System. Requiere una migración explícita.")
    if type(document["format_version"]) is not int or document["format_version"] != BACKUP_VERSION:
        raise ValueError("Versión de respaldo no compatible.")
    if type(document["schema_version"]) is not int or document["schema_version"] != SCHEMA_VERSION:
        raise ValueError("Versión de base de datos no compatible. No se modificó tu base.")
    digest = document["sha256"]
    unsigned = {key: value for key, value in document.items() if key != "sha256"}
    if not isinstance(digest, str) or digest != _hash(unsigned):
        raise ValueError("El respaldo fue alterado o está dañado: su checksum no coincide.")
    if not isinstance(document["created_at"], str) or not document["created_at"].endswith("Z") or not isinstance(document["app_version"], str):
        raise ValueError("Los metadatos del respaldo no son válidos.")
    try:
        datetime.fromisoformat(document["created_at"].replace("Z", "+00:00"))
    except ValueError:
        raise ValueError("La fecha del respaldo no es válida.") from None
    payload = document["payload"]
    if not isinstance(payload, dict) or set(payload) != {"tables"} or not isinstance(payload["tables"], dict):
        raise ValueError("El respaldo no contiene sus tablas.")
    if set(payload["tables"]) != set(TABLES) or not isinstance(document["counts"], dict) or set(document["counts"]) != set(TABLES):
        raise ValueError("El respaldo contiene tablas faltantes o desconocidas.")
    for name in TABLES:
        rows = payload["tables"][name]
        if not isinstance(rows, list) or type(document["counts"][name]) is not int or document["counts"][name] != len(rows):
            raise ValueError(f"El recuento de la tabla {name} no coincide.")
    return payload["tables"]


def _insert_tables(db: Database, tables: dict):
    for name in TABLES:
        columns = [row["name"] for row in db.connection.execute(f'PRAGMA table_info("{name}")')]
        column_set = set(columns)
        sql = f'INSERT INTO "{name}" ({",".join(chr(34) + column + chr(34) for column in columns)}) VALUES ({",".join("?" for _ in columns)})'
        for row in tables[name]:
            if not isinstance(row, dict) or set(row) != column_set:
                raise ValueError(f"Campos faltantes o desconocidos en {name}.")
            for key, value in row.items():
                if value is not None and type(value) not in (str, int):
                    raise ValueError(f"Tipo inválido en {name}.{key}.")
                if type(value) is int and abs(value) > MAX_CENTS:
                    raise ValueError(f"Número fuera de rango en {name}.{key}.")
                if key in ("business_date", "completed_date", "cancelled_date") and value is not None and _day(value) != value:
                    raise ValueError(f"Fecha inválida en {name}.{key}.")
            db.connection.execute(sql, [row[column] for column in columns])


def validate_database(db: Database):
    """Check cross-table invariants not expressible as SQLite CHECK constraints."""
    conn = db.connection
    if conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
        raise ValueError("La base no superó el control de integridad.")
    if conn.execute("PRAGMA foreign_key_check").fetchone() is not None:
        raise ValueError("Hay referencias inexistentes en el respaldo.")
    # Existing backups without currency metadata retain the legacy precision.
    # A PYG marker promises whole guaranies in every financial record, not only
    # in its on-screen presentation, so reject hidden fractional balances.
    quantum = money_quantum(db)
    if quantum > 1:
        for table in TABLES:
            columns = [row["name"] for row in conn.execute(f'PRAGMA table_info("{table}")')
                       if row["name"].endswith("_cents")]
            if not columns:
                continue
            for row in conn.execute(f'SELECT {",".join(columns)} FROM "{table}"'):
                if any(row[column] is not None and row[column] % quantum for column in columns):
                    raise ValueError(f"El respaldo PYG contiene fracciones de guaraní en {table}.")
        if conn.execute("SELECT 1 FROM orders WHERE discount_kind='fixed' AND discount_value % ? != 0 LIMIT 1", (quantum,)).fetchone():
            raise ValueError("El respaldo PYG contiene un descuento fijo con fracciones de guaraní.")
    products = {}
    for product in conn.execute("SELECT * FROM products"):
        if product["normalized_name"] != _normal(product["name"]) or product["normalized_name"] in products:
            raise ValueError("El modelo normalizado no coincide o está duplicado.")
        products[product["normalized_name"]] = product["id"]
    unique_skus = set()
    unique_variants = set()
    for variant in conn.execute("SELECT * FROM variants"):
        sku = _normal(variant["sku"])
        key = (variant["product_id"], _normal(variant["size"]), _normal(variant["color"]))
        if sku in unique_skus or key in unique_variants:
            raise ValueError("Hay SKU o combinaciones de talle/color duplicadas.")
        unique_skus.add(sku)
        unique_variants.add(key)
    variants = {row["id"] for row in conn.execute("SELECT id FROM variants")}
    balances = {row["variant_id"]: dict(row) for row in conn.execute("SELECT * FROM inventory")}
    if variants != set(balances):
        raise ValueError("Cada variante debe tener exactamente un saldo de inventario.")
    for variant_id, balance in balances.items():
        movements = conn.execute("SELECT * FROM inventory_movements WHERE variant_id=? ORDER BY sequence", (variant_id,)).fetchall()
        qty = reserved = value = 0
        for movement in movements:
            qty += movement["quantity_delta"]
            reserved += movement["reserved_delta"]
            value += movement["value_delta_cents"]
            if qty < 0 or reserved < 0 or reserved > qty or value < 0 or (qty == 0 and value != 0):
                raise ValueError("El historial de stock contiene una posición inválida.")
        if (qty, reserved, value) != (balance["on_hand"], balance["reserved"], balance["value_cents"]):
            raise ValueError("El saldo de inventario no coincide con sus movimientos.")
        active = conn.execute("SELECT COALESCE(SUM(quantity),0) FROM reservations WHERE variant_id=? AND state='active'", (variant_id,)).fetchone()[0]
        if active != reserved:
            raise ValueError("Las reservas activas no coinciden con el inventario reservado.")
    for reservation in conn.execute("SELECT r.*,o.status,i.order_id AS item_order,i.variant_id AS item_variant,i.quantity AS item_qty FROM reservations r JOIN orders o ON o.id=r.order_id JOIN order_items i ON i.id=r.order_item_id"):
        if reservation["item_order"] != reservation["order_id"] or reservation["item_variant"] != reservation["variant_id"] or reservation["quantity"] != reservation["item_qty"]:
            raise ValueError("Una reserva no corresponde a la línea del pedido.")
        if reservation["state"] == "active" and reservation["status"] != "reserved":
            raise ValueError("Una reserva activa requiere un pedido reservado.")
    for order in conn.execute("SELECT * FROM orders"):
        items = conn.execute("SELECT * FROM order_items WHERE order_id=?", (order["id"],)).fetchall()
        if not items or sum(item["gross_cents"] for item in items) != order["subtotal_cents"] or sum(item["discount_cents"] for item in items) != order["discount_cents"] or sum(item["net_cents"] for item in items) != order["total_cents"]:
            raise ValueError("Los importes del pedido no coinciden con sus artículos.")
        for item in items:
            if item["quantity"] * item["unit_price_cents"] != item["gross_cents"]:
                raise ValueError("El precio por cantidad no coincide con el detalle.")
            returned = conn.execute("SELECT COALESCE(SUM(ri.quantity),0),COALESCE(SUM(ri.revenue_cents),0),COALESCE(SUM(ri.cogs_cents),0) FROM return_items ri JOIN returns r ON r.id=ri.return_id WHERE ri.order_item_id=? AND r.order_id=?", (item["id"], order["id"])).fetchone()
            if returned[0] > item["quantity"] or returned[1] > item["net_cents"] or (item["cogs_cents"] is not None and returned[2] > item["cogs_cents"]):
                raise ValueError("Una devolución supera lo vendido.")
        known = all(item["cogs_cents"] is not None for item in items)
        if order["completed_date"] is not None:
            if (known and order["cogs_cents"] != sum(item["cogs_cents"] for item in items)) or (not known and order["cogs_cents"] is not None):
                raise ValueError("El CMV del pedido no coincide con sus líneas.")
        elif order["cogs_cents"] is not None or order["status"] == "completed":
            raise ValueError("Un pedido sin completar no puede tener CMV reconocido.")
        if order["status"] == "reserved":
            active_count = conn.execute("SELECT COUNT(*) FROM reservations WHERE order_id=? AND state='active'", (order["id"],)).fetchone()[0]
            if active_count != len(items):
                raise ValueError("El pedido reservado no tiene todas sus líneas reservadas.")
        receipts = refunds = 0
        for payment in conn.execute("SELECT * FROM payments WHERE order_id=?", (order["id"],)):
            if payment["business_date"] < order["business_date"]:
                raise ValueError("Un cobro no puede preceder al pedido.")
            if payment["kind"] == "receipt":
                receipts += payment["amount_cents"]
                if round_money_ratio(db, payment["amount_cents"] * payment["commission_bps"], 10000) != payment["commission_cents"]:
                    raise ValueError("La comisión del cobro no coincide con su base y porcentaje.")
            else:
                refunds += payment["amount_cents"]
                if payment["commission_cents"] != 0:
                    raise ValueError("Un reembolso no revierte automáticamente comisiones.")
        if refunds > receipts:
            raise ValueError("Los reembolsos superan los cobros registrados.")
        returns = conn.execute("SELECT COALESCE(SUM(revenue_cents),0),COALESCE(SUM(cogs_cents),0) FROM returns WHERE order_id=?", (order["id"],)).fetchone()
        financial = {bucket: sum(row[0] for row in conn.execute("SELECT amount_cents FROM economic_entries WHERE order_id=? AND bucket=?", (order["id"], bucket))) for bucket in ("revenue", "cogs")}
        expected_revenue = order["total_cents"] - returns[0] if order["completed_date"] else 0
        expected_cogs = sum(item["cogs_cents"] or 0 for item in items) - returns[1] if order["completed_date"] else 0
        if financial["revenue"] != expected_revenue or financial["cogs"] != expected_cogs:
            raise ValueError("Los hechos económicos no coinciden con la venta y sus devoluciones.")
    for returned in conn.execute("SELECT * FROM returns"):
        items = conn.execute("SELECT ri.*,oi.order_id FROM return_items ri JOIN order_items oi ON oi.id=ri.order_item_id WHERE ri.return_id=?", (returned["id"],)).fetchall()
        if not items or any(item["order_id"] != returned["order_id"] for item in items) or sum(item["revenue_cents"] for item in items) != returned["revenue_cents"] or sum(item["cogs_cents"] for item in items) != returned["cogs_cents"]:
            raise ValueError("La devolución y sus detalles no coinciden.")
    for draft in conn.execute("SELECT payload FROM drafts"):
        try:
            decoded = json.loads(draft[0], parse_constant=lambda value: (_ for _ in ()).throw(ValueError(f"Número no finito: {value}")))
            if not isinstance(decoded, dict):
                raise ValueError
        except (json.JSONDecodeError, ValueError):
            raise ValueError("Hay un borrador guardado con formato inválido.") from None
    _validate_financial_ledgers(db)


def _validate_financial_ledgers(db: Database):
    """Reconcile every recognized cost and cash movement with its source record.

    A valid checksum alone must not allow an altered commission or bank balance.
    Recognition dates are exact; later settlement dates remain legitimate for
    accrued operating expenses and direct shipping/packaging costs.
    """
    conn = db.connection
    expected_economic = Counter()
    expected_cash = {}

    def economic(day, bucket, amount, order_id=None, channel=None, reference=""):
        if amount:
            expected_economic[(day, bucket, amount, order_id, channel, reference)] += 1

    def cash(day, kind, reference, method, amount, order_id=None, exact_date=True):
        if amount:
            key = (kind, reference)
            if key in expected_cash:
                raise ValueError("Dos documentos pretenden el mismo movimiento de caja.")
            expected_cash[key] = (day, method, amount, order_id, exact_date)

    for order in conn.execute("SELECT * FROM orders"):
        if order["completed_date"]:
            economic(order["completed_date"], "revenue", order["total_cents"], order["id"], order["channel"], order["id"])
            cogs = sum(row[0] or 0 for row in conn.execute("SELECT cogs_cents FROM order_items WHERE order_id=?", (order["id"],)))
            economic(order["completed_date"], "cogs", cogs, order["id"], order["channel"], order["id"])
            for bucket in ("shipping", "packaging"):
                amount = order[f"{bucket}_cents"]
                economic(order["completed_date"], bucket, amount, order["id"], order["channel"], order["id"])
                if order["direct_costs_paid"]:
                    cash(order["completed_date"], bucket, order["id"], order["direct_cost_method"], -amount, order["id"], False)
    channels = {row["id"]: row["channel"] for row in conn.execute("SELECT id,channel FROM orders")}
    for payment in conn.execute("SELECT * FROM payments"):
        amount = payment["amount_cents"] if payment["kind"] == "receipt" else -payment["amount_cents"]
        cash(payment["business_date"], payment["kind"], payment["id"], payment["method"], amount, payment["order_id"])
        if payment["commission_cents"]:
            economic(payment["business_date"], "commission", payment["commission_cents"], payment["order_id"], channels[payment["order_id"]], payment["id"])
            cash(payment["business_date"], "commission", payment["id"], payment["method"], -payment["commission_cents"], payment["order_id"])
    for returned in conn.execute("SELECT * FROM returns"):
        economic(returned["business_date"], "revenue", -returned["revenue_cents"], returned["order_id"], channels[returned["order_id"]], returned["id"])
        economic(returned["business_date"], "cogs", -returned["cogs_cents"], returned["order_id"], channels[returned["order_id"]], returned["id"])
    for expense in conn.execute("SELECT * FROM expenses"):
        economic(expense["business_date"], "operating_expense", expense["amount_cents"], channel=expense["channel"], reference=expense["id"])
        if expense["paid"]:
            cash(expense["business_date"], "expense", expense["id"], expense["payment_method"], -expense["amount_cents"], exact_date=False)
    for purchase in conn.execute("SELECT * FROM purchases"):
        if purchase["paid"]:
            cash(purchase["business_date"], "purchase", purchase["id"], purchase["method"], -purchase["total_cost_cents"])
    for movement in conn.execute("SELECT * FROM inventory_movements WHERE kind='adjustment' AND quantity_delta<0"):
        economic(movement["business_date"], "inventory_loss", -movement["value_delta_cents"], reference=movement["variant_id"])

    actual_economic = Counter((row["business_date"], row["bucket"], row["amount_cents"], row["order_id"], row["channel"], row["reference"])
                              for row in conn.execute("SELECT * FROM economic_entries"))
    if actual_economic != expected_economic:
        raise ValueError("Los costos, comisiones o fechas del libro económico no coinciden con sus documentos.")
    for row in conn.execute("SELECT * FROM cash_entries"):
        origin = expected_cash.pop((row["kind"], row["reference"]), None)
        if origin is None:
            raise ValueError("La caja contiene un movimiento duplicado o sin documento de origen.")
        day, method, amount, order_id, exact = origin
        if (row["method"], row["delta_cents"], row["order_id"]) != (method, amount, order_id):
            raise ValueError("Un importe o medio de caja no coincide con su cobro o gasto.")
        invalid_date = row["business_date"] != day if exact else row["business_date"] < day
        if invalid_date:
            raise ValueError("La fecha de caja no coincide con el movimiento que la originó.")
    if expected_cash:
        raise ValueError("Faltan movimientos de caja correspondientes a cobros o gastos pagados.")


def validate_backup(document: dict) -> dict:
    tables = _validate_envelope(document)
    staged = Database(":memory:")
    try:
        with staged.transaction():
            staged.connection.execute("PRAGMA defer_foreign_keys=ON")
            _insert_tables(staged, tables)
            validate_database(staged)
    except sqlite3.Error as error:
        raise ValueError(f"El respaldo viola una restricción de la base de datos: {error}") from error
    finally:
        staged.close()
    return tables


def restore_json(db: Database, path: str | Path) -> dict:
    tables = validate_backup(_load_backup(path))
    folder = Path(db.path).parent if db.path != ":memory:" else Path(path).parent
    stamp = datetime.now().strftime("%Y-%m-%d-%H%M%S-%f")
    previous = folder / "backups" / f"antes-de-restaurar-{stamp}.json"
    # Preserve the existing database before the destructive transaction starts.
    _atomic_file(previous, json.dumps(make_backup(db), ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8"))
    try:
        with db.transaction():
            db.connection.execute("PRAGMA defer_foreign_keys=ON")
            for name in reversed(TABLES):
                db.connection.execute(f'DELETE FROM "{name}"')
            _insert_tables(db, tables)
            validate_database(db)
    except sqlite3.Error as error:
        raise ValueError(f"No se restauró la base; se conservan los datos anteriores: {error}") from error
    return {"counts": {name: len(tables[name]) for name in TABLES}, "previous_backup": str(previous)}


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "1" if value else "0"
    return str(value)


def read_import_file(path: str | Path, worksheet: str | None = None) -> tuple[list[str], list[dict[str, str]]]:
    path = Path(path)
    if path.stat().st_size > MAX_IMPORT_BYTES:
        raise ValueError("El archivo supera 10 MB. Dividilo en archivos más pequeños.")
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeError:
            raise ValueError("El CSV debe estar guardado como UTF-8.") from None
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = []
        try:
            for row in csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True):
                if len(raw) > MAX_ROWS:
                    raise ValueError(f"El archivo supera {MAX_ROWS} filas de datos.")
                raw.append(row)
        except csv.Error as error:
            raise ValueError(f"El CSV contiene comillas o filas inválidas: {error}") from error
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        import zipfile
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("El Excel descomprimido supera el límite permitido.")
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if worksheet is not None and worksheet not in book.sheetnames:
                raise ValueError("La hoja seleccionada no existe.")
            sheet = book[worksheet] if worksheet else book.worksheets[0]
            raw = []
            for row in sheet.iter_rows():
                if len(raw) > MAX_ROWS:
                    raise ValueError(f"El archivo supera {MAX_ROWS} filas de datos.")
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("El archivo contiene fórmulas. Pegá sus valores antes de importar.")
                raw.append([_cell_text(cell.value) for cell in row])
        finally:
            book.close()
    else:
        raise ValueError("Seleccioná un archivo .csv o .xlsx; no se admiten macros ni .xls antiguos.")
    raw = [row for row in raw if any(str(value).strip() for value in row)]
    if not raw:
        raise ValueError("El archivo está vacío.")
    if len(raw) > MAX_ROWS + 1:
        raise ValueError(f"El archivo supera {MAX_ROWS} filas de datos.")
    headers = [str(value).strip() for value in raw[0]]
    if any(not value for value in headers) or len({_normal(value) for value in headers}) != len(headers):
        raise ValueError("Los encabezados deben ser únicos y no pueden estar vacíos.")
    records = []
    for number, row in enumerate(raw[1:], 2):
        if len(row) > len(headers) and any(str(value).strip() for value in row[len(headers):]):
            raise ValueError(f"Fila {number}: hay datos sin encabezado.")
        records.append({header: _cell_text(row[index]) if index < len(row) else "" for index, header in enumerate(headers)})
    return headers, records


@dataclass
class ImportPreview:
    kind: str
    headers: list[str]
    raw_rows: list[dict[str, str]]
    mapping: dict[str, str]
    source_account: str
    records: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    content_hash: str = ""

    @property
    def ready(self) -> bool:
        return bool(self.records) and not self.errors


def _whole(value: str, label: str, minimum=0) -> int:
    if not str(value).strip().isdigit():
        raise ValueError(f"{label}: ingresá unidades enteras.")
    return integer(int(value), label, minimum, 1_000_000_000)


def _import_money(db: Database, value: str) -> int:
    """Read PYG whole units, including Paraguay's dot thousands grouping.

    Decimal-zero Excel values are accepted, but nonzero fractional guaranies
    are rejected instead of rounding imported payments or changing their value.
    ARS legacy files retain the previous explicit decimal parser.
    """
    if currency_code(db) != "PYG":
        return parse_money(value)
    text = str(value).strip()
    text = re.sub(r"^(?:₲\s*|Gs\.?\s*)", "", text, flags=re.IGNORECASE)
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,0{1,2})?", text):
        whole = text.split(",")[0].replace(".", "")
    elif re.fullmatch(r"\d+(?:[.,]0{1,2})?", text):
        whole = re.split(r"[.,]", text)[0]
    else:
        raise ValueError("Ingresá guaraníes enteros, por ejemplo 50000 o 50.000, sin decimales.")
    return validate_amount(db, int(whole) * 100, "Importe en guaraníes")


def _customer_id(service, row: dict) -> tuple[str | None, str | None]:
    customer_id = row.get("cliente_id", "").strip()
    name = row.get("cliente", "").strip()
    customers = service.list_customers()
    if customer_id:
        if not any(customer["id"] == customer_id for customer in customers):
            raise ValueError("El cliente_id no existe en esta base.")
        return customer_id, None
    if not name:
        return None, None
    matches = [customer for customer in customers if _normal(customer["name"]) == _normal(name)]
    if len(matches) > 1:
        raise ValueError("Hay varias clientas con ese nombre; indicá cliente_id para resolverlo.")
    return (matches[0]["id"], None) if matches else (None, name)


def _build_preview(service, preview: ImportPreview) -> ImportPreview:
    preview.records = []
    preview.errors = []
    preview.warnings = []
    if preview.kind not in REQUIRED_COLUMNS:
        raise ValueError("El tipo de importación no es válido.")
    missing = REQUIRED_COLUMNS[preview.kind] - {name for name, source in preview.mapping.items() if source}
    if missing:
        preview.errors.append("Faltan columnas obligatorias: " + ", ".join(sorted(missing)))
        return preview
    if any(source not in preview.headers for source in preview.mapping.values() if source):
        preview.errors.append("El mapeo refiere a una columna inexistente.")
        return preview
    variants = service.list_variants()
    sku_map = {_normal(row["sku"]): row for row in variants}
    file_skus = set()
    file_variants = set()
    existing_variants = {(_normal(variant["model"]), _normal(variant["size"]), _normal(variant["color"])) for variant in variants}
    groups = {}
    for number, raw in enumerate(preview.raw_rows, 2):
        row = {name: raw.get(source, "").strip() for name, source in preview.mapping.items() if source}
        try:
            if preview.kind == "variants":
                if not all(row.get(key) for key in ("modelo", "sku", "talle", "color")):
                    raise ValueError("Modelo, SKU, talle y color son obligatorios.")
                sku = _normal(row["sku"])
                if sku in sku_map or sku in file_skus:
                    raise ValueError("El SKU ya existe en la base o se repite en el archivo.")
                combination = (_normal(row["modelo"]), _normal(row["talle"]), _normal(row["color"]))
                if combination in existing_variants or combination in file_variants:
                    raise ValueError("La combinación modelo+talle+color ya existe o se repite.")
                file_skus.add(sku)
                file_variants.add(combination)
                preview.records.append({
                    "model": row["modelo"], "category": row.get("categoria") or "Lencería", "sku": row["sku"],
                    "size": row["talle"], "color": row["color"], "retail_cents": _import_money(service.db, row["pvp"]),
                    "wholesale_cents": _import_money(service.db, row["mayorista"]), "min_stock": _whole(row.get("stock_minimo") or "1", "Stock mínimo"),
                    "initial_qty": _whole(row["stock"], "Stock"), "initial_unit_cost_cents": _import_money(service.db, row["costo"]),
                })
            else:
                external_id = row["id_externo"].strip()
                if not external_id:
                    raise ValueError("id_externo es obligatorio para evitar ventas duplicadas.")
                key = "treinta:" + _canonical([preview.source_account, external_id]).decode("utf-8")
                if service.db.connection.execute("SELECT 1 FROM orders WHERE external_key=?", (key,)).fetchone():
                    raise ValueError("La venta externa ya fue importada en esta cuenta.")
                variant = sku_map.get(_normal(row["sku"]))
                if variant is None:
                    raise ValueError("El SKU no existe; importá primero el catálogo.")
                channel = next((value for value in CHANNELS if _normal(value) == _normal(row["canal"])), None)
                if channel is None:
                    raise ValueError("Canal inválido: usá Instagram, WhatsApp, TikTok o Local.")
                method = next((value for value in (*PAYMENT_METHODS, "Crédito") if _normal(value) == _normal(row["medio_pago"])), None)
                if method is None:
                    raise ValueError("Falta un medio de cobro conocido. No se inventan cobros al importar.")
                customer_id, customer_name = _customer_id(service, row)
                if method == "Crédito" and not (customer_id or customer_name):
                    raise ValueError("Una venta fiada requiere cliente o cliente_id.")
                sale_day = _day(row["fecha"])
                payment_day = _day(row.get("fecha_cobro") or sale_day)
                if payment_day < sale_day:
                    raise ValueError("La fecha de cobro no puede preceder a la venta.")
                common = {
                    "external_key": key, "business_date": sale_day, "channel": channel, "payment_method": method,
                    "customer_id": customer_id, "customer_name": customer_name, "payment_date": payment_day,
                    "shipping_cents": _import_money(service.db, row.get("envio") or "0"), "packaging_cents": _import_money(service.db, row.get("packaging") or "0"),
                    "discount_cents": _import_money(service.db, row.get("descuento") or "0"), "commission_bps": percent_bps(row.get("comision_porcentaje") or "0"),
                }
                line = {"variant_id": variant["id"], "model": variant["model"], "variant": f'{variant["size"]} · {variant["color"]}',
                        "quantity": _whole(row["cantidad"], "Cantidad", 1), "unit_price_cents": _import_money(service.db, row["precio_unitario"]),
                        "unit_cost_cents": _import_money(service.db, row["costo_unitario"]) if row.get("costo_unitario") else None}
                integer(line["quantity"] * line["unit_price_cents"], "Subtotal")
                if line["unit_cost_cents"] is not None:
                    integer(line["quantity"] * line["unit_cost_cents"], "Costo histórico")
                group = groups.get(key)
                if group is None:
                    groups[key] = {**common, "items": [line], "row": number}
                else:
                    if any(group[field] != value for field, value in common.items()):
                        raise ValueError("Las filas de una misma venta tienen fecha, cliente, descuento o costos distintos. Los importes de pedido deben repetirse iguales, no dividirse.")
                    same = next((item for item in group["items"] if item["variant_id"] == line["variant_id"]), None)
                    if same:
                        if (same["unit_price_cents"], same["unit_cost_cents"]) != (line["unit_price_cents"], line["unit_cost_cents"]):
                            raise ValueError("Un SKU repetido en el pedido tiene precios/costos diferentes; conciliá esas líneas antes de importar.")
                        same["quantity"] += line["quantity"]
                        integer(same["quantity"], "Cantidad total", 1, 1_000_000_000)
                    else:
                        group["items"].append(line)
        except (ValueError, KeyError) as error:
            preview.errors.append(f"Fila {number}: {error}")
    if preview.kind == "sales":
        preview.records = list(groups.values())
        for group in preview.records:
            subtotal = sum(item["quantity"] * item["unit_price_cents"] for item in group["items"])
            if group["discount_cents"] > subtotal:
                preview.errors.append(f'Venta de fila {group["row"]}: el descuento supera el subtotal.')
            if any(item["unit_cost_cents"] is None for item in group["items"]):
                preview.warnings.append(f'Venta de fila {group["row"]}: falta costo histórico; la ganancia real quedará incompleta, no se usará costo cero.')
        preview.warnings.append("Ventas históricas: no cambian el stock actual. En medios distintos de Crédito se registran cobros efectivos según la columna medio_pago y fecha_cobro; conciliá esos datos antes de confirmar.")
        preview.warnings.append("Envío, packaging, descuento y comisión se leen una vez por documento, aunque estén repetidos en cada artículo. Los campos opcionales vacíos se consideran cero, previa confirmación.")
    return preview


def preview_import(service, path: str | Path, kind: str, mapping: dict[str, str] | None = None, source_account="principal") -> ImportPreview:
    if not str(source_account).strip():
        raise ValueError("Indicá una cuenta de origen para detectar duplicados.")
    headers, rows = read_import_file(path)
    canonical = VARIANT_COLUMNS if kind == "variants" else SALE_COLUMNS
    if mapping is None:
        mapping = {name: next((header for header in headers if _normal(header) == _normal(name)), "") for name in canonical}
    preview = ImportPreview(kind, headers, rows, mapping, str(source_account).strip())
    preview.content_hash = _hash({"rows": rows, "mapping": mapping, "source_account": preview.source_account, "kind": kind})
    return _build_preview(service, preview)


def _insert(conn, table: str, **values):
    conn.execute(f'INSERT INTO "{table}" ({",".join(values)}) VALUES ({",".join("?" for _ in values)})', tuple(values.values()))


def _next_sequence(conn, table: str) -> int:
    return conn.execute(f'SELECT COALESCE(MAX(sequence),0)+1 FROM "{table}"').fetchone()[0]


def _import_sale(service, group: dict) -> str:
    conn = service.db.connection
    order_id = str(uuid4())
    customer_id = group["customer_id"]
    if group["customer_name"]:
        matches = [customer for customer in service.list_customers() if _normal(customer["name"]) == _normal(group["customer_name"])]
        customer_id = matches[0]["id"] if len(matches) == 1 else service.save_customer(group["customer_name"])
    gross = [item["quantity"] * item["unit_price_cents"] for item in group["items"]]
    discounts = allocate_money(service.db, group["discount_cents"], gross)
    costs = [item["quantity"] * item["unit_cost_cents"] if item["unit_cost_cents"] is not None else None for item in group["items"]]
    cogs = sum(costs) if all(cost is not None for cost in costs) else None
    subtotal = integer(sum(gross), "Subtotal")
    total = subtotal - group["discount_cents"]
    terms = "credit" if group["payment_method"] == "Crédito" else "immediate"
    _insert(conn, "orders", id=order_id, number=f"IMP-{uuid4().hex[:10].upper()}", business_date=group["business_date"], completed_date=group["business_date"],
            channel=group["channel"], customer_id=customer_id, status="completed", terms=terms, payment_method=group["payment_method"],
            subtotal_cents=subtotal, discount_kind="fixed", discount_value=group["discount_cents"], discount_cents=group["discount_cents"], total_cents=total,
            cogs_cents=cogs, commission_bps=group["commission_bps"], shipping_cents=group["shipping_cents"], packaging_cents=group["packaging_cents"],
            direct_cost_method="Efectivo", direct_costs_paid=0, source="import", external_key=group["external_key"], created_at=utc_now(), version=0,
            notes="Importación histórica: inventario actual sin cambios. Costos directos reconocidos; su pago requiere conciliación.")
    for index, item in enumerate(group["items"]):
        _insert(conn, "order_items", id=str(uuid4()), order_id=order_id, variant_id=item["variant_id"], model_snapshot=item["model"], variant_snapshot=item["variant"],
                quantity=item["quantity"], unit_price_cents=item["unit_price_cents"], gross_cents=gross[index], discount_cents=discounts[index],
                net_cents=gross[index] - discounts[index], cogs_cents=costs[index])
    known_cogs = sum(cost for cost in costs if cost is not None)
    for bucket, amount in (("revenue", total), ("cogs", known_cogs), ("shipping", group["shipping_cents"]), ("packaging", group["packaging_cents"])):
        if amount is not None and amount:
            _insert(conn, "economic_entries", id=str(uuid4()), business_date=group["business_date"], bucket=bucket, amount_cents=amount,
                    order_id=order_id, channel=group["channel"], reference=order_id)
    if terms == "immediate" and total:
        payment_id = str(uuid4())
        fee = round_money_ratio(service.db, total * group["commission_bps"], 10000)
        _insert(conn, "payments", id=payment_id, order_id=order_id, business_date=group["payment_date"], amount_cents=total,
                method=group["payment_method"], kind="receipt", commission_bps=group["commission_bps"], commission_cents=fee, created_at=utc_now())
        _insert(conn, "cash_entries", id=str(uuid4()), sequence=_next_sequence(conn, "cash_entries"), business_date=group["payment_date"], method=group["payment_method"],
                delta_cents=total, kind="receipt", order_id=order_id, reference=payment_id, note="Cobro histórico importado")
        if fee:
            _insert(conn, "economic_entries", id=str(uuid4()), business_date=group["payment_date"], bucket="commission", amount_cents=fee,
                    order_id=order_id, channel=group["channel"], reference=payment_id)
            _insert(conn, "cash_entries", id=str(uuid4()), sequence=_next_sequence(conn, "cash_entries"), business_date=group["payment_date"], method=group["payment_method"],
                    delta_cents=-fee, kind="commission", order_id=order_id, reference=payment_id, note="Comisión histórica importada")
    return order_id


def apply_import(service, preview: ImportPreview, acknowledge_warnings=False) -> dict:
    if not isinstance(preview, ImportPreview):
        raise ValueError("La vista previa de importación no es válida.")
    digest = _hash({"rows": preview.raw_rows, "mapping": preview.mapping, "source_account": preview.source_account, "kind": preview.kind})
    if digest != preview.content_hash:
        raise ValueError("La vista previa cambió. Generala de nuevo antes de importar.")
    with service.db.transaction():
        _build_preview(service, preview)
        if preview.errors or not preview.records:
            raise ValueError("No se puede importar: " + (preview.errors[0] if preview.errors else "no hay registros válidos."))
        if preview.warnings and not acknowledge_warnings:
            raise ValueError("Confirmá las advertencias de la vista previa antes de importar históricos.")
        identifiers = []
        for record in preview.records:
            identifiers.append(service.save_variant(record) if preview.kind == "variants" else _import_sale(service, record))
        validate_database(service.db)
        service.db.set_meta("last_import", json.dumps({"kind": preview.kind, "source_account": preview.source_account, "documents": len(identifiers), "at": utc_now()}))
    return {"count": len(identifiers), "ids": identifiers}


def _safe_csv_text(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _money_decimal(cents, currency="ARS"):
    if cents is None:
        return ""
    sign = "-" if cents < 0 else ""
    whole, fraction = divmod(abs(cents), 100)
    if currency == "PYG" and fraction == 0:
        return Decimal(f"{sign}{whole}")
    return Decimal(f"{sign}{whole}.{fraction:02d}")


def export_table(service, path: str | Path, kind: str) -> Path:
    path = Path(path)
    _protect_database_path(service.db, path)
    currency = currency_code(service.db)
    monetary = lambda amount: _money_decimal(amount, currency)
    if kind == "variants":
        headers = list(VARIANT_COLUMNS)
        rows = [[v["model"], v["category"], v["sku"], v["size"], v["color"], monetary(v["cpp_cents"]), monetary(v["retail_cents"]),
                 monetary(v["wholesale_cents"]), v["on_hand"], v["min_stock"]] for v in service.list_variants()]
    elif kind == "sales":
        headers = ["numero", "fecha", "estado", "cliente", "canal", "modelo", "variante", "cantidad", "precio_unitario", "costo_total_linea", "descuento_linea", "venta_neta_linea", "envio_pedido", "packaging_pedido", "saldo_pedido"]
        rows = []
        for order in service.list_orders():
            full = service.get_order(order["id"])
            for index, item in enumerate(full["items"]):
                rows.append([full["number"], full["completed_date"] or full["business_date"], full["status"], full.get("customer_name") or "", full["channel"], item["model_snapshot"], item["variant_snapshot"], item["quantity"],
                             monetary(item["unit_price_cents"]), monetary(item["cogs_cents"]), monetary(item["discount_cents"]), monetary(item["net_cents"]),
                             monetary(full["shipping_cents"]) if index == 0 else "", monetary(full["packaging_cents"]) if index == 0 else "", monetary(full["balance_cents"]) if index == 0 else ""])
    elif kind == "expenses":
        headers = ["fecha", "categoria", "tipo", "descripcion", "importe", "medio_pago", "canal", "pagado"]
        rows = [[row["business_date"], row["category"], row["cost_type"], row["description"], monetary(row["amount_cents"]), row["payment_method"], row["channel"] or "", row["paid"]]
                for row in service.db.connection.execute("SELECT * FROM expenses ORDER BY business_date,id")]
    else:
        raise ValueError("Elegí variantes, ventas o gastos.")
    if path.suffix.casefold() == ".csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(headers)
        writer.writerows([[_safe_csv_text(value) for value in row] for row in rows])
        return _atomic_file(path, ("\ufeff" + stream.getvalue()).encode("utf-8"))
    if path.suffix.casefold() != ".xlsx":
        raise ValueError("La exportación debe terminar en .csv o .xlsx.")
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Sou Linda System"
    sheet.append(headers)
    for values in rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif isinstance(cell.value, Decimal):
                cell.number_format = '"₲" #,##0' if currency == "PYG" else "#,##0.00"
    book.properties.description = f"Moneda del negocio: {currency}. Importes expresados en unidades monetarias, sin conversión de divisas."
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(45, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return _atomic_file(path, stream.getvalue())
